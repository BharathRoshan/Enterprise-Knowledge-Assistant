import os
from pathlib import Path

import gradio as gr
from langchain_community.document_loaders import PyPDFLoader, TextLoader
from langchain_community.vectorstores import FAISS
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter
from transformers import pipeline

EMBED_MODEL = os.getenv("EMBED_MODEL", "BAAI/bge-small-en-v1.5")
GEN_MODEL = os.getenv("GEN_MODEL", "google/flan-t5-small")
CHUNK_SIZE = 800
CHUNK_OVERLAP = 120
TOP_K = 4

print("Loading embedding model...")
embedding_model = HuggingFaceEmbeddings(
    model_name=EMBED_MODEL,
    encode_kwargs={"normalize_embeddings": True},
)

print("Loading answer model...")
generator = pipeline("text2text-generation", model=GEN_MODEL)


def load_document(path: str):
    """Load supported document formats without deserializing user-supplied objects."""
    file_path = Path(path)
    suffix = file_path.suffix.lower()

    if suffix == ".pdf":
        return PyPDFLoader(str(file_path)).load()
    if suffix in {".txt", ".md"}:
        return TextLoader(str(file_path), encoding="utf-8", autodetect_encoding=True).load()
    if suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("Install python-docx to read DOCX files.") from exc
        document = Document(str(file_path))
        text = "\n".join(p.text for p in document.paragraphs if p.text.strip())
        from langchain_core.documents import Document as LCDocument
        return [LCDocument(page_content=text, metadata={"source": file_path.name})] if text else []

    raise ValueError(f"Unsupported file type: {suffix or 'no extension'}")


def build_knowledge_base(files):
    if not files:
        return None, "Upload one or more PDF, DOCX, TXT, or Markdown files first."

    documents = []
    errors = []
    for file in files:
        path = file if isinstance(file, str) else getattr(file, "name", None)
        if not path:
            errors.append("Skipped a file with no readable path.")
            continue
        try:
            loaded = load_document(path)
            for doc in loaded:
                doc.metadata["source"] = Path(path).name
            documents.extend(loaded)
        except Exception as exc:
            errors.append(f"{Path(path).name}: {exc}")

    documents = [doc for doc in documents if doc.page_content and doc.page_content.strip()]
    if not documents:
        details = "\n".join(errors)
        return None, "No readable text was found. Check the files and try again.\n" + details

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=CHUNK_SIZE,
        chunk_overlap=CHUNK_OVERLAP,
        add_start_index=True,
    )
    chunks = splitter.split_documents(documents)
    if not chunks:
        return None, "The uploaded files did not contain text that could be indexed."

    store = FAISS.from_documents(chunks, embedding_model)
    message = f"Knowledge base ready: {len(documents)} document page(s) split into {len(chunks)} searchable chunks."
    if errors:
        message += "\nSome files could not be loaded:\n" + "\n".join(errors)
    return store, message


def answer_question(question, store):
    if not question or not question.strip():
        return "Enter a question to get started.", "No sources retrieved."
    if store is None:
        return "Upload documents and click **Build knowledge base** first.", "No sources retrieved."

    docs = store.similarity_search(question.strip(), k=TOP_K)
    if not docs:
        return "I could not find relevant information in the uploaded documents.", "No sources retrieved."

    context_blocks = []
    for index, doc in enumerate(docs, start=1):
        source = doc.metadata.get("source", "Unknown file")
        page = doc.metadata.get("page")
        page_text = f", page {int(page) + 1}" if isinstance(page, (int, float)) else ""
        context_blocks.append(f"[Source {index}: {source}{page_text}]\n{doc.page_content}")
    context = "\n\n".join(context_blocks)[:5000]

    prompt = (
        "You are an enterprise knowledge assistant. Answer the question using only the context. "
        "If the context does not contain the answer, say you could not find it in the uploaded documents. "
        "Do not invent facts. Be concise and include source names when useful.\n\n"
        f"Context:\n{context}\n\nQuestion: {question.strip()}\nAnswer:"
    )
    try:
        result = generator(prompt, max_new_tokens=160, do_sample=False, truncation=True)
        answer = result[0]["generated_text"].strip()
    except Exception as exc:
        return f"Generation failed: {exc}", "Retrieved context is available, but no answer was generated."

    source_lines = []
    for index, doc in enumerate(docs, start=1):
        source = doc.metadata.get("source", "Unknown file")
        page = doc.metadata.get("page")
        page_text = f" · Page {int(page) + 1}" if isinstance(page, (int, float)) else ""
        snippet = " ".join(doc.page_content.split())[:280]
        source_lines.append(f"**{index}. {source}{page_text}**\n\n> {snippet}")
    return answer, "\n\n---\n\n".join(source_lines)


with gr.Blocks(theme=gr.themes.Soft(), title="Enterprise Knowledge Assistant") as demo:
    gr.Markdown(
        "# Enterprise Knowledge Assistant\n"
        "Upload trusted business documents, build a searchable knowledge base, and ask questions with source snippets."
    )
    knowledge_base = gr.State(None)

    with gr.Row():
        files_input = gr.File(
            label="Documents",
            file_count="multiple",
            file_types=[".pdf", ".docx", ".txt", ".md"],
            type="filepath",
        )
        with gr.Column():
            build_button = gr.Button("Build knowledge base", variant="primary")
            status_output = gr.Markdown("Upload documents to begin.")

    question_input = gr.Textbox(
        label="Your question",
        placeholder="What do the documents say about customer due diligence?",
        lines=3,
    )
    ask_button = gr.Button("Ask", variant="primary")
    with gr.Row():
        answer_output = gr.Markdown(label="Answer")
        sources_output = gr.Markdown(label="Retrieved sources")

    build_button.click(
        fn=build_knowledge_base,
        inputs=files_input,
        outputs=[knowledge_base, status_output],
    )
    ask_button.click(
        fn=answer_question,
        inputs=[question_input, knowledge_base],
        outputs=[answer_output, sources_output],
    )
    question_input.submit(
        fn=answer_question,
        inputs=[question_input, knowledge_base],
        outputs=[answer_output, sources_output],
    )

if __name__ == "__main__":
    demo.queue().launch()
